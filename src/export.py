import re
from io import BytesIO

from docx import Document

from src.schema import Minutes


def safe_filename(title: str) -> str:
    """Turn a meeting title into a safe file name (no extension)."""
    name = re.sub(r"[^\w\- ]+", "", title).strip().replace(" ", "_")
    return name[:60] or "meeting_minutes"


# ---------------- Markdown ----------------

def _cell(text) -> str:
    return (text or "-").replace("|", "\\|").replace("\n", " ")


def minutes_to_markdown(m: Minutes) -> str:
    out: list[str] = [f"# {m.title}", ""]
    if m.date:
        out += [f"**Date:** {m.date}", ""]
    if m.attendees:
        out += [f"**Attendees:** {', '.join(m.attendees)}", ""]

    def bullets(heading, items, empty=None):
        if not items and empty is None:
            return
        out.extend([f"## {heading}", ""])
        if items:
            out.extend(f"- {item}" for item in items)
        else:
            out.append(empty)
        out.append("")

    bullets("Agenda", m.agenda)
    if m.summary:
        out.extend(["## Summary", "", m.summary, ""])
    bullets("Discussion Points", m.discussion_points)
    bullets("Decisions", m.decisions, empty="_None recorded._")

    out.extend(["## Action Items", ""])
    if m.action_items:
        out.extend(["| Task | Owner | Due |", "|---|---|---|"])
        for a in m.action_items:
            out.append(f"| {_cell(a.task)} | {_cell(a.owner)} | {_cell(a.due)} |")
    else:
        out.append("_None recorded._")
    out.append("")

    bullets("Open Questions", m.open_questions)
    if m.next_meeting:
        out.extend(["## Next Meeting", "", m.next_meeting, ""])
    return "\n".join(out).rstrip() + "\n"


# ---------------- Word ----------------

def minutes_to_docx(m: Minutes) -> bytes:
    doc = Document()
    doc.add_heading(m.title, level=0)

    def meta(label, value):
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    if m.date:
        meta("Date", m.date)
    if m.attendees:
        meta("Attendees", ", ".join(m.attendees))

    def bullets(heading, items, empty=None):
        if not items and empty is None:
            return
        doc.add_heading(heading, level=1)
        if items:
            for item in items:
                doc.add_paragraph(item, style="List Bullet")
        else:
            doc.add_paragraph(empty)

    bullets("Agenda", m.agenda)
    if m.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(m.summary)
    bullets("Discussion Points", m.discussion_points)
    bullets("Decisions", m.decisions, empty="None recorded.")

    doc.add_heading("Action Items", level=1)
    if m.action_items:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["Task", "Owner", "Due"]):
            cell.text = ""
            cell.paragraphs[0].add_run(text).bold = True
        for a in m.action_items:
            row = table.add_row().cells
            row[0].text = a.task
            row[1].text = a.owner or "-"
            row[2].text = a.due or "-"
    else:
        doc.add_paragraph("None recorded.")

    bullets("Open Questions", m.open_questions)
    if m.next_meeting:
        doc.add_heading("Next Meeting", level=1)
        doc.add_paragraph(m.next_meeting)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()