from io import BytesIO

from docx import Document

from src.export import minutes_to_docx, minutes_to_markdown, safe_filename
from src.schema import ActionItem, Minutes

M = Minutes(
    title="Project Sync",
    date="14 October 2026",
    attendees=["Priya", "Rahul"],
    summary="We met.",
    decisions=["Move the server"],
    action_items=[ActionItem(task="Send palette | final", owner="Meera", due="Friday")],
    next_meeting="21 October",
)


def test_markdown_has_sections_and_table():
    md = minutes_to_markdown(M)
    assert md.startswith("# Project Sync")
    assert "**Attendees:** Priya, Rahul" in md
    assert "- Move the server" in md
    assert "| Send palette \\| final | Meera | Friday |" in md  # pipe is escaped
    assert "## Next Meeting" in md


def test_markdown_empty_sections():
    md = minutes_to_markdown(Minutes())
    assert "_None recorded._" in md
    assert "## Agenda" not in md


def test_docx_roundtrip():
    doc = Document(BytesIO(minutes_to_docx(M)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Project Sync" in text and "Move the server" in text
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Task", "Owner", "Due"]
    assert table.rows[1].cells[1].text == "Meera"


def test_safe_filename():
    assert safe_filename("Q4 Plan / Review: 2026!") == "Q4_Plan__Review_2026"
    assert safe_filename("???") == "meeting_minutes"